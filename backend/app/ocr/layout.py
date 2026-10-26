"""From what the OCR model read to a Word file laid out the way the page was: blocks become paragraphs with the page's alignment and
weight, HTML tables become Word tables. The viewer, the parser and the index all read that file, so tables stay tables there."""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import Path
from typing import Any

from docx import Document as WordDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

_MD_PREFIX = re.compile(r"^\s{0,3}(#{1,6}\s+|>\s+)")
_MD_EMPHASIS = re.compile(r"(\*\*|__|`|(?<=\s)#{1,6}\s)")  # also a heading mark in the middle of a line (blocks joined)
_TABLE_RULE = re.compile(r"^\s*\|?\s*:?-{2,}:?\s*(\|\s*:?-{2,}:?\s*)*\|?\s*$")
_TABLE_HTML = re.compile(r"<table\b.*?</table>", re.IGNORECASE | re.DOTALL)
_SKIPPED = {"image", "header_image", "footer_image", "page_number", "chart"}  # stamps, pictures and page numbers carry no text of the document
_ALIGN = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT, "justify": WD_ALIGN_PARAGRAPH.JUSTIFY}


@dataclass
class Para:
    text: str
    align: str = "left"
    bold: bool = False
    italic: bool = False
    size: float | None = None
    src: int | None = None  # index of the block it came from


@dataclass
class Table:
    rows: list[list[tuple[str, int]]] = field(default_factory=list)  # (cell text, columns it spans)
    src: int | None = None


@dataclass
class Band:
    """Blocks that sat side by side on the page (the two columns of a letterhead, "Số" and the date): a table without borders."""

    cells: list[list[Para]]
    widths: list[float]


class _TableReader(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.rows: list[list[tuple[str, int]]] = []
        self._row: list[tuple[str, int]] | None = None
        self._cell: list[str] | None = None
        self._span = 1

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag == "tr":
            self._row = []
        elif tag in ("td", "th"):
            self._cell = []
            span = dict(attrs).get("colspan") or "1"
            self._span = int(span) if span.isdigit() and 0 < int(span) <= 50 else 1
        elif tag == "br" and self._cell is not None:
            self._cell.append(" ")

    def handle_endtag(self, tag: str) -> None:
        if tag in ("td", "th") and self._cell is not None and self._row is not None:
            self._row.append((" ".join("".join(self._cell).split()), self._span))
            self._cell = None
        elif tag == "tr" and self._row is not None:
            if self._row:
                self.rows.append(self._row)
            self._row = None

    def handle_data(self, data: str) -> None:
        if self._cell is not None:
            self._cell.append(data)


def read_table(html: str) -> Table:
    reader = _TableReader()
    reader.feed(html)
    reader.close()
    return Table(reader.rows)


def clean(line: str) -> str:
    return _MD_EMPHASIS.sub("", _MD_PREFIX.sub("", line)).strip()


def _markdown_table(lines: list[str]) -> Table:
    rows = []
    for line in lines:
        if _TABLE_RULE.match(line):
            continue
        rows.append([(c.strip(), 1) for c in line.strip().strip("|").split("|")])
    return Table(rows)


_LIST_START = re.compile(r"^\s*([-+•*]|\d+[.)]|[a-zđ][.)])\s")
_DATE_LINE = re.compile(r"^.{2,40},\s*ngày\s+.*\bnăm\b", re.IGNORECASE)
_LINE_HEIGHT = 17  # height of one line of text on the 0-1000 scale of a page, when the page gives no better idea


def _upper_ratio(text: str) -> float:
    letters = [c for c in text if c.isalpha()]
    return sum(c.isupper() for c in letters) / len(letters) if letters else 0.0


def _alignment(box: tuple[int, int, int, int] | list[int] | None, label: str, text: str = "") -> str:
    """Left, centred, right or justified from where the block sat on the page (box on a 0-1000 scale)."""
    if not box:
        return "center" if label == "title" else "left"
    x0, _, x1, _ = box
    centre, width = (x0 + x1) / 2, x1 - x0
    listed = bool(_LIST_START.match(text)) or label == "list"
    if not listed and abs(centre - 500) <= 90 and (width <= 700 or label == "title"):  # a full-width paragraph is body text, not a centred line
        return "center"
    if not listed and x0 >= 450 and width <= 520:
        return "right"
    if label == "text" and width >= 650 and len(text) > 120:
        return "justify"
    return "left"


def _lines_of(text: str, box: list[int] | tuple[int, ...] | None, line_height: float) -> list[str]:
    """The lines of a header-like block. The model joins the lines of a block with spaces, but the box says how many there were."""
    parts = [p for p in text.split("\n") if p.strip()]
    if len(parts) > 1 or not box:
        return parts
    wanted = max(1, round((box[3] - box[1]) / line_height))
    one = parts[0] if parts else ""
    one = re.sub(r"\s+(Độc lập\s*-\s*Tự do\s*-\s*Hạnh phúc)\s*$", r"\n\1", one)  # the motto always sits under the national title
    if "\n" in one:
        return [p for p in one.split("\n") if p.strip()]
    words = one.split()
    if wanted <= 1 or len(words) < wanted:
        return [one]
    target, lines, current = len(one) / wanted, [], []
    for word in words:
        if current and len(" ".join(current + [word])) > target * 1.15 and len(lines) < wanted - 1:
            lines.append(" ".join(current))
            current = []
        current.append(word)
    lines.append(" ".join(current))
    return lines


def _text_items(text: str) -> list[Para | Table]:
    """Plain Markdown text of a page (no blocks): lines, with pipe tables kept as tables."""
    out: list[Para | Table] = []
    pipe: list[str] = []
    for raw in text.splitlines():
        if raw.strip().startswith("|") and raw.strip().endswith("|"):
            pipe.append(raw)
            continue
        if pipe:
            out.append(_markdown_table(pipe))
            pipe = []
        line = clean(raw)
        bold = raw.strip().startswith(("**", "#"))
        out.append(Para(line, bold=bold))
    if pipe:
        out.append(_markdown_table(pipe))
    return out


def _split_html(text: str) -> list[Para | Table]:
    """Text that may carry HTML tables: the tables become tables, the rest lines."""
    out: list[Para | Table] = []
    last = 0
    for match in _TABLE_HTML.finditer(text):
        out += _text_items(text[last:match.start()])
        out.append(read_table(match.group()))
        last = match.end()
    return out + _text_items(text[last:])


def _line_height(blocks: list[dict[str, Any]]) -> float:
    """Height of one text line on this page: the middle of the heights of its short text blocks."""
    heights = sorted(b["box"][3] - b["box"][1] for b in blocks if b["label"] == "text" and len(b["text"]) < 90 and "\n" not in b["text"] and b.get("box"))
    return float(heights[len(heights) // 2]) if len(heights) >= 3 and heights[len(heights) // 2] >= 8 else float(_LINE_HEIGHT)


def _side_by_side(a: dict[str, Any], b: dict[str, Any]) -> bool:
    """Whether two blocks stand next to each other: they share most of their height and do not overlap horizontally."""
    if not a.get("box") or not b.get("box"):
        return False
    ay0, ay1, by0, by1 = a["box"][1], a["box"][3], b["box"][1], b["box"][3]
    shared = min(ay1, by1) - max(ay0, by0)
    return shared > 0.5 * min(ay1 - ay0, by1 - by0) and (a["box"][2] <= b["box"][0] or b["box"][2] <= a["box"][0])


_TEXTUAL = {"text", "title", "header", "caption", "list", "footnote", "footer", "formula"}


def _paragraphs(block: dict[str, Any], line_height: float, src: int | None = None) -> list[Para]:
    label, body = block["label"], block["text"]
    align = _alignment(block.get("box"), label, body)
    header_like = label in ("title", "header") or align == "center"
    raws = _lines_of(body, block.get("box"), line_height) if header_like and label != "list" else body.split("\n")
    out = []
    for raw in raws:
        line = clean(raw)
        if not line:
            continue
        bold = raw.strip().startswith("**") or (label in ("title", "header") and (align == "center" or _upper_ratio(line) > 0.6))
        out.append(Para(line, align=align, bold=bold,
                        italic=label == "footnote" or bool(_DATE_LINE.match(line)), size=10 if label == "footnote" else None, src=src))
    return out


def page_items(text: str, blocks: list[dict[str, Any]] | None) -> list[Para | Table | Band]:
    """What to write for one page: from its blocks when the model gave them, else from its text."""
    if not blocks:
        return _split_html(text)
    out: list[Para | Table | Band] = []
    line_height = _line_height(blocks)
    shown = [(n, b) for n, b in enumerate(blocks) if b["label"] not in _SKIPPED]
    i = 0
    while i < len(shown):
        at, block = shown[i]
        if block["label"] == "table" or _TABLE_HTML.search(block["text"]):
            for item in _split_html(block["text"]):
                if isinstance(item, Para) and not item.text:
                    continue
                item.src = at
                out.append(item)
            i += 1
            continue
        row = [shown[i]]
        while (i + len(row) < len(shown) and shown[i + len(row)][1]["label"] in _TEXTUAL and block["label"] in _TEXTUAL
               and all(_side_by_side(m[1], shown[i + len(row)][1]) for m in row)):
            row.append(shown[i + len(row)])
        if len(row) > 1:
            row.sort(key=lambda m: m[1]["box"][0])
            cells = [_paragraphs(b, line_height, n) for n, b in row]
            widths = [max(1, b["box"][2] - b["box"][0]) for _, b in row]
            out.append(Band([[Para(p.text, "center" if p.align == "right" else p.align, p.bold, p.italic, p.size, p.src) for p in cell] for cell in cells], widths))
        else:
            out += _paragraphs(block, line_height, at)
        i += len(row)
    return out


def flat_texts(items: list[Para | Table | Band]) -> list[tuple[int | None, str]]:
    """``(block index, text)`` of every non-empty paragraph in the order the Word file holds them (cells row by row, a band cell by cell)."""
    out: list[tuple[int | None, str]] = []
    for item in items:
        if isinstance(item, Para):
            out.append((item.src, item.text))
        elif isinstance(item, Table):
            out += [(item.src, text) for row in item.rows for text, _ in row if text]
        else:
            out += [(p.src, p.text) for cell in item.cells for p in cell]
    return [(src, text) for src, text in out if text.strip()]


def write_docx(pages: list[tuple[str, list[dict[str, Any]] | None]], target: Path) -> None:
    """Writes the pages (``(text, blocks)`` each) as one Word file, a page break between pages."""
    doc = WordDocument()
    for number, (text, blocks) in enumerate(pages):
        if number:
            doc.add_page_break()
        for item in page_items(text, blocks):
            if isinstance(item, Table):
                _add_table(doc, item)
                continue
            if isinstance(item, Band):
                _add_band(doc, item)
                continue
            paragraph = doc.add_paragraph()
            paragraph.alignment = _ALIGN[item.align]
            run = paragraph.add_run(item.text)
            run.bold, run.italic = item.bold or None, item.italic or None
            if item.size:
                run.font.size = Pt(item.size)
    doc.save(str(target))


def _add_band(doc: Any, band: Band) -> None:
    """Side-by-side blocks as a one-row table without borders (python-docx's default table style has none)."""
    word = doc.add_table(rows=1, cols=len(band.cells))
    word.autofit = False
    total = float(sum(band.widths))
    for cell, paragraphs, width in zip(word.rows[0].cells, band.cells, band.widths):
        cell.width = Cm(16.0 * width / total)
        first = True
        for item in paragraphs:
            paragraph = cell.paragraphs[0] if first else cell.add_paragraph()
            first = False
            paragraph.alignment = _ALIGN[item.align]
            run = paragraph.add_run(item.text)
            run.bold, run.italic = item.bold or None, item.italic or None


def _add_table(doc: Any, table: Table) -> None:
    columns = max((sum(span for _, span in row) for row in table.rows), default=0)
    if not columns:
        return
    word = doc.add_table(rows=len(table.rows), cols=columns)
    word.style = "Table Grid"
    for r, row in enumerate(table.rows):
        at = 0
        for text, span in row:
            cell = word.cell(r, at)
            if span > 1:
                cell = cell.merge(word.cell(r, min(at + span - 1, columns - 1)))
            cell.text = text
            at += span
