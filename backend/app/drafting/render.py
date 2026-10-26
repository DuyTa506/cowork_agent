"""Builds the .docx of a draft.
"""
from __future__ import annotations

import itertools
import re
from collections.abc import Callable, Iterator
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from . import metrics, tables
from .content import DraftContent, blank, split

FONT = "Times New Roman"
BODY_PT, HEADER_PT, SMALL_PT = 14, 13, 12
RECIPIENT_PT = 11
LINE_HEIGHT = 1.15  # one line of Times New Roman at single spacing, in font sizes
_VV = re.compile(r"^v/v\s*:?\s*", re.IGNORECASE)  # a trích yếu the user typed with its "V/v" already
HEAD_CM, FOOT_CM = (6.0, 10.0), (6.5, 9.5)  # header and signature columns; together the 16 cm text width
CELL_MARGIN_CM = 0.1
ISSUER_RULE, TITLE_RULE, KEM_THEO_RULE = 0.42, 0.33, 0.4  # line length as a share of the text above it
_ALIGN = {"left": WD_ALIGN_PARAGRAPH.LEFT, "center": WD_ALIGN_PARAGRAPH.CENTER, "right": WD_ALIGN_PARAGRAPH.RIGHT}
_LEAD = re.compile(r"^(Điều\s+\d+[a-z]?\.)(?=\s)")
_KHOAN_NUMBER = re.compile(r"^\d{1,2}\.(?=\s)")
_NAMESPACES = ('xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" '
               'xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing" '
               'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
               'xmlns:wps="http://schemas.microsoft.com/office/word/2010/wordprocessingShape"')
_LINE_SHAPE = (
    '<w:r {ns}><w:drawing><wp:anchor distT="0" distB="0" distL="0" distR="0" simplePos="0" relativeHeight="{id}" behindDoc="0" '
    'locked="0" layoutInCell="1" allowOverlap="1"><wp:simplePos x="0" y="0"/><wp:positionH relativeFrom="column"><wp:align>center'
    '</wp:align></wp:positionH><wp:positionV relativeFrom="paragraph"><wp:posOffset>{y}</wp:posOffset></wp:positionV>'
    '<wp:extent cx="{cx}" cy="0"/><wp:effectExtent l="0" t="0" r="0" b="0"/><wp:wrapNone/><wp:docPr id="{id}" name="Line {id}"/>'
    '<wp:cNvGraphicFramePr/><a:graphic><a:graphicData uri="http://schemas.microsoft.com/office/word/2010/wordprocessingShape">'
    '<wps:wsp><wps:cNvCnPr/><wps:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="{cx}" cy="0"/></a:xfrm><a:prstGeom prst="line">'
    '<a:avLst/></a:prstGeom><a:ln w="6350"><a:solidFill><a:srgbClr val="000000"/></a:solidFill></a:ln></wps:spPr><wps:bodyPr/>'
    '</wps:wsp></a:graphicData></a:graphic></wp:anchor></w:drawing></w:r>')


def _run(paragraph, text: str, size: float, *, bold: bool = False, italic: bool = False):
    """Adds a run in the document font.

    The font is set for East-Asian text too, which Word otherwise renders in another face.

    Args:
        paragraph: The paragraph to add to.
        text: Text of the run.
        size: Font size in points.
        bold: Bold.
        italic: Italic.

    Returns:
        The new run.
    """
    run = paragraph.add_run(text)
    run.font.name, run.font.size, run.bold, run.italic = FONT, Pt(size), bold, italic
    properties = run._element.get_or_add_rPr()
    fonts = properties.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        properties.append(fonts)
    for attribute in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        fonts.set(qn(attribute), FONT)
    return run


def _format(paragraph, align=WD_ALIGN_PARAGRAPH.LEFT, *, before: float = 0, after: float = 0, first_cm: float = 0, keep: bool = False):
    """Sets the geometry of a paragraph.

    Args:
        paragraph: The paragraph.
        align: Alignment.
        before: Space before, in points.
        after: Space after, in points.
        first_cm: First-line indent in cm (0 = none).
        keep: Keep with the next paragraph.

    Returns:
        The paragraph, at single line spacing.
    """
    layout = paragraph.paragraph_format
    paragraph.alignment = align
    layout.space_before, layout.space_after = Pt(before), Pt(after)
    layout.line_spacing_rule, layout.first_line_indent = WD_LINE_SPACING.SINGLE, Cm(first_cm) if first_cm else None
    layout.keep_with_next = keep
    return paragraph


def _line(cell_or_doc, text: str, size: float, align=WD_ALIGN_PARAGRAPH.CENTER, *, bold=False, italic=False, before=0, after=0,
          first_cm=0, keep=False, reuse=None):
    """Adds a formatted paragraph holding one run.

    Args:
        cell_or_doc: The document or table cell to add to.
        text: The text; an empty text gives an empty paragraph.
        size: Font size in points.
        align: Alignment.
        bold: Bold.
        italic: Italic.
        before: Space before, in points.
        after: Space after, in points.
        first_cm: First-line indent in cm.
        keep: Keep with the next paragraph.
        reuse: An existing empty paragraph to fill instead (a cell's first one).

    Returns:
        The paragraph.
    """
    paragraph = reuse if reuse is not None else cell_or_doc.add_paragraph()
    _format(paragraph, align, before=before, after=after, first_cm=first_cm, keep=keep)
    if text:
        _run(paragraph, text, size, bold=bold, italic=italic)
    return paragraph


def _rule(paragraph, ids: Iterator[int], length_cm: float, *, lines: int, size: float) -> None:
    """Anchors a horizontal line to a paragraph, centred in its column, just below its text.

    Args:
        paragraph: The paragraph the line belongs to.
        ids: Source of drawing ids, unique within the document.
        length_cm: Length of the line in cm.
        lines: Number of lines the paragraph's text takes.
        size: Font size of that text in points.
    """
    below_pt = lines * size * LINE_HEIGHT + 2
    shape = _LINE_SHAPE.format(ns=_NAMESPACES, id=next(ids), y=round(below_pt * 12700), cx=round(length_cm * 360000))
    paragraph._p.append(parse_xml(shape))


def _ruled(cell_or_doc, text: str, size: float, room_cm: float, share: float, ids: Iterator[int], *, bold: bool = False,
           italic: bool = False, before: float = 0, reuse=None, keep: bool = False):
    """Adds a centred paragraph with a line under it.

    Args:
        cell_or_doc: The document or table cell to add to.
        text: The text.
        size: Font size in points.
        room_cm: Width the text wraps in, in cm.
        share: Length of the line as a share of the longest wrapped line (capped at `room_cm`).
        ids: Source of drawing ids.
        bold: Bold.
        italic: Italic.
        before: Space before, in points.
        reuse: An existing empty paragraph to fill instead.
        keep: Keep with the next paragraph.

    Returns:
        The paragraph.
    """
    paragraph = _line(cell_or_doc, text, size, bold=bold, italic=italic, before=before, reuse=reuse, keep=keep)
    wrapped = metrics.wrap(text, size, room_cm, bold=bold) or [text]
    longest = max(metrics.width_cm(line, size, bold=bold) for line in wrapped)
    _rule(paragraph, ids, min(longest * share, room_cm), lines=len(wrapped), size=size)
    return paragraph


def _borderless(table) -> None:
    """Removes every border of a layout table and narrows its cell margins (children inserted in schema order)."""
    properties = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), "nil")
        borders.append(element)
    properties.insert_element_before(borders, "w:shd", "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption",
                                     "w:tblDescription", "w:tblPrChange")
    margins = OxmlElement("w:tblCellMar")
    for edge in ("left", "right"):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:w"), str(round(CELL_MARGIN_CM * 567)))
        element.set(qn("w:type"), "dxa")
        margins.append(element)
    properties.insert_element_before(margins, "w:tblLook", "w:tblCaption", "w:tblDescription", "w:tblPrChange")


def _columns(table, widths_cm) -> None:
    """Fixes the column widths of a table (python-docx needs both the grid and every cell set)."""
    table.autofit = False
    for column, width in zip(table.columns, widths_cm):
        column.width = Cm(width)
        for cell in column.cells:
            cell.width = Cm(width)


def _lines(text: str) -> list[str]:
    """The non-empty lines of `text`, stripped."""
    return [line.strip() for line in text.splitlines() if line.strip()]


def _placeholder(text: str, fallback: str) -> str:
    """The field text, or `fallback` while it is still blank, so the gap is visible on the page."""
    return text.strip() if not blank(text) else fallback


def _closed(text: str) -> str:
    """The last paragraph of the body, ending with the closing mark "./." that the circular prescribes."""
    return text if text.endswith("./.") else text.rstrip(" .;,") + "./."


def _recipients(text: str) -> list[str]:
    """The "Kính gửi" addressees as printed.

    Args:
        text: The field, one addressee per line (a leading "Kính gửi:" or dash is ignored).

    Returns:
        One addressee ending with "."; or several, each on a dash line ending with ";", the last with ".".
    """
    names = [re.sub(r"^\s*(kính gửi\s*:)?\s*[-–+•]?\s*", "", line, flags=re.IGNORECASE).rstrip(" .;,") for line in _lines(text)]
    names = [n for n in names if n] or ["…"]
    if len(names) == 1:
        return [f"{names[0]}."]
    return [f"- {name}{'.' if i == len(names) - 1 else ';'}" for i, name in enumerate(names)]


def _kinh_gui(document, text: str, room_cm: float = 16.0) -> None:
    """Adds "Kính gửi:" with its addressees (Thông tư 68, Phụ lục I mục II.9).

    One addressee sits on a centred line. Several form a block centred on the page, each dash aligned under the first one after
    the colon.

    Args:
        document: The document.
        text: The "Kính gửi" field.
        room_cm: Text width of the page in cm.
    """
    lines = _recipients(text)
    label = "Kính gửi: "
    if len(lines) == 1:
        _line(document, label + lines[0], BODY_PT, before=12)
        return
    label_cm = metrics.width_cm(label, BODY_PT)
    block = label_cm + max(metrics.width_cm(line, BODY_PT) for line in lines)
    left = max((room_cm - block) / 2, 0)
    for index, line in enumerate(lines):
        paragraph = _format(document.add_paragraph(), WD_ALIGN_PARAGRAPH.LEFT, before=12 if index == 0 else 0)
        layout = paragraph.paragraph_format
        layout.left_indent = Cm(left + label_cm)
        if index == 0:
            layout.first_line_indent = Cm(-label_cm)
            _run(paragraph, label, BODY_PT)
        _run(paragraph, line, BODY_PT)


def _body_line(document, block: dict[str, Any], text: str, *, articles: bool) -> None:
    """Adds one paragraph of the body, bold as its numbering level calls for (Thông tư 68, Phụ lục I mục II.8).

    Args:
        document: The document.
        block: The block (its kind decides the emphasis).
        text: The text to print (the block's, with the closing mark when it is the last one).
        articles: Whether the document is divided into "Điều"; the number of a khoản inside an article stays regular.
    """
    if block["kind"] == "heading":
        _line(document, text, BODY_PT, WD_ALIGN_PARAGRAPH.JUSTIFY, bold=True, before=6, first_cm=1.0, keep=True)
        return
    paragraph = _format(document.add_paragraph(), WD_ALIGN_PARAGRAPH.JUSTIFY, before=6, first_cm=1.0)
    lead = _LEAD.match(text) if block["kind"] == "para" else None
    if lead is None and block["kind"] == "para" and not articles:
        lead = _KHOAN_NUMBER.match(text)
    if lead:
        _run(paragraph, lead.group(0), BODY_PT, bold=True)
        text = text[lead.end():]
    _run(paragraph, text, BODY_PT)


def _page_numbers(section) -> None:
    """Puts the page number centred in the top margin of a section, 14 pt, hidden on its first page."""
    section.different_first_page_header_footer = True
    paragraph = _format(section.header.paragraphs[0], WD_ALIGN_PARAGRAPH.CENTER)
    run = _run(paragraph, "", BODY_PT)
    for kind, text in (("begin", None), (None, " PAGE "), ("separate", None), (None, "2"), ("end", None)):
        if kind:
            element = OxmlElement("w:fldChar")
            element.set(qn("w:fldCharType"), kind)
        else:
            element = OxmlElement("w:instrText" if text.strip() == "PAGE" else "w:t")
            element.set(qn("xml:space"), "preserve")
            element.text = text
        run._r.append(element)


def _page(section, *, landscape: bool) -> float:
    """A4 with the Thông tư 68 margins (the same when turned landscape); returns the text width in cm."""
    section.orientation = WD_ORIENT.LANDSCAPE if landscape else WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = (Cm(29.7), Cm(21)) if landscape else (Cm(21), Cm(29.7))
    section.top_margin = section.bottom_margin = Cm(2.0)
    section.left_margin, section.right_margin = Cm(3.0), Cm(2.0)
    return tables.LANDSCAPE_CM if landscape else tables.PORTRAIT_CM


def _restart_numbering(section) -> None:
    """Numbers the pages of a section from 1, hiding the number on its first page (an appendix, Mẫu 15a)."""
    properties = section._sectPr
    numbering = properties.find(qn("w:pgNumType"))
    if numbering is None:
        numbering = OxmlElement("w:pgNumType")
        properties.insert_element_before(numbering, "w:cols", "w:formProt", "w:vAlign", "w:noEndnote", "w:titlePg",
                                         "w:textDirection", "w:bidi", "w:rtlGutter", "w:docGrid", "w:printerSettings",
                                         "w:sectPrChange")
    numbering.set(qn("w:start"), "1")
    section.different_first_page_header_footer = True


def _keep_row(row, *, repeat: bool = False) -> None:
    """Keeps a table row from breaking across pages.

    Args:
        row: The row.
        repeat: A heading row, printed again at the top of each page.
    """
    properties = row._tr.get_or_add_trPr()
    properties.append(OxmlElement("w:cantSplit"))
    if repeat:
        properties.append(OxmlElement("w:tblHeader"))


def _cell_text(cell, text: str, align, *, bold: bool = False, italic_parens: bool = False) -> None:
    """Writes the text of a table cell, one paragraph per line, vertically centred.

    Args:
        cell: The cell.
        text: The text.
        align: Horizontal alignment.
        bold: Bold.
        italic_parens: Print the part from "(" italic and not bold, as a section row does ("Bổ sung (…)").
    """
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    for index, line in enumerate(text.splitlines() or [""]):
        paragraph = _format(cell.paragraphs[0] if index == 0 else cell.add_paragraph(), align, before=2, after=2)
        if italic_parens and "(" in line:
            head, _, rest = line.partition("(")
            _run(paragraph, head, tables.TABLE_PT, bold=bold)
            _run(paragraph, f"({rest}", tables.TABLE_PT, italic=True)
        elif line:
            _run(paragraph, line, tables.TABLE_PT, bold=bold)


def _table(document, spec: dict[str, Any], room_cm: float) -> None:
    """Adds a table of the draft.

    Grid lines; the heading rows bold, merged as the columns group, and repeated on each page; group and total rows bold; a
    section row merged across the table and centred.

    Args:
        document: The document.
        spec: A normalised table spec (see `tables.normalize`).
        room_cm: Text width of the page in cm; the columns fill it.
    """
    heading = tables.header(spec)
    width = len(tables.leaves(spec))
    table = document.add_table(rows=len(heading) + len(spec["rows"]), cols=width)
    table.style = document.styles["Table Grid"]
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _columns(table, tables.widths(spec, room_cm))
    for r, cells in enumerate(heading):
        for item in cells:
            cell = table.cell(r, item["col"])
            if item["colspan"] > 1:
                cell = cell.merge(table.cell(r, item["col"] + item["colspan"] - 1))
            if item["rowspan"] > 1:
                cell = cell.merge(table.cell(r + item["rowspan"] - 1, item["col"]))
            _cell_text(cell, item["text"], WD_ALIGN_PARAGRAPH.CENTER, bold=True)
        _keep_row(table.rows[r], repeat=True)
    aligns = tables.aligns(spec)
    for offset, row in enumerate(spec["rows"]):
        r = len(heading) + offset
        _keep_row(table.rows[r])
        if row["type"] == "section":
            cell = table.cell(r, 0).merge(table.cell(r, width - 1)) if width > 1 else table.cell(r, 0)
            _cell_text(cell, row["cells"][0], WD_ALIGN_PARAGRAPH.CENTER, bold=True, italic_parens=True)
            continue
        bold = row["type"] in ("group", "total")
        for c, text in enumerate(row["cells"]):
            _cell_text(table.cell(r, c), text, _ALIGN[aligns[c]], bold=bold)
    _line(document, "", SMALL_PT)  # Word needs a paragraph between a table and what follows


def _blocks(document, blocks: list[dict[str, Any]], room_cm: float, *, articles: bool, close: bool) -> None:
    """Adds a run of body blocks.

    Args:
        document: The document.
        blocks: The blocks, in order.
        room_cm: Text width of the page in cm (for tables).
        articles: Whether the document is divided into "Điều" (see `_body_line`).
        close: End the last paragraph or item with "./." (the main body, not an appendix).
    """
    last = max((i for i, b in enumerate(blocks) if b["kind"] in ("para", "item")), default=-1) if close else -1
    for index, block in enumerate(blocks):
        if block["kind"] == "table":
            if block.get("table"):
                _table(document, block["table"], room_cm)
            continue
        text = _closed(block["text"]) if index == last else block["text"]
        _body_line(document, block, text, articles=articles)


def render(content: DraftContent, target: Path, *, for_check: bool = False, issuer_name: str = "") -> Path:
    """Writes a draft as a .docx.

    Args:
        content: The draft.
        target: Path of the file to write; its folder is created when missing.
        for_check: Replace the parts the registry clerk and the drafter fill in after signing (the number in "Số: …/TM-KHQS", the
            drafter's initials on the "Lưu" line) with well-formed samples, so the format check judges the layout rather than
            reporting them as malformed.
        issuer_name: The issuing body as an appendix cites it ("của Bộ Tham mưu"); without it the header's name is used.

    Returns:
        `target`.
    """
    template, fields = content.template, content.fields
    ids = itertools.count(1)

    def text(key: str, fallback: str = "…") -> str:
        """A field as printed: `fallback` while blank, samples in place of the clerk's and drafter's parts with `for_check`."""
        value = _placeholder(fields[key]["text"], fallback) if key in fields else ""
        if for_check and key == "so_ky_hieu" and "…" in value:
            numbered = value.replace("…", "01", 1)
            return numbered if "…" not in numbered else f"Số: 01/{template.kind}-ĐV"
        if for_check and key == "noi_nhan":
            return value.replace("…", "A01")
        return value

    document = Document()
    style = document.styles["Normal"]
    style.font.name, style.font.size = FONT, Pt(BODY_PT)
    style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    section = document.sections[0]
    room = _page(section, landscape=False)
    _page_numbers(section)

    # ----- header: issuing body, number (and a công văn's subject) | national name, motto, place and date -----
    header = document.add_table(rows=1, cols=2)
    header.alignment = WD_TABLE_ALIGNMENT.CENTER
    _borderless(header)
    _columns(header, HEAD_CM)
    left, right = header.rows[0].cells
    left_room, right_room = (w - 2 * CELL_MARGIN_CM for w in HEAD_CM)
    first = left.paragraphs[0]
    owner_line = text("co_quan_chu_quan", "")
    if owner_line:
        _line(left, owner_line, HEADER_PT, reuse=first)
        first = None
    _ruled(left, text("co_quan_ban_hanh"), HEADER_PT, left_room, ISSUER_RULE, ids, bold=True, reuse=first)
    _line(left, text("so_ky_hieu", "Số: …/…"), HEADER_PT, before=6)
    is_letter = not text("ten_loai", "") and "trich_yeu" in fields
    if is_letter:  # Công văn: "V/v ..." centred under the number, 12 pt upright (ô 5b)
        _line(left, f"V/v {_VV.sub('', text('trich_yeu'))}", SMALL_PT, before=6)
    _line(right, text("quoc_hieu"), HEADER_PT, bold=True, reuse=right.paragraphs[0])
    _ruled(right, text("tieu_ngu"), HEADER_PT, right_room, 1.0, ids, bold=True)
    if "dia_danh_ngay" in fields:
        _line(right, text("dia_danh_ngay", "……, ngày … tháng … năm …"), HEADER_PT, italic=True, before=6)

    # ----- title (ô 5a) and addressees -----
    kind_word = text("ten_loai", "")
    if kind_word:
        _line(document, kind_word, BODY_PT, bold=True, before=12, keep=True)
        _ruled(document, text("trich_yeu"), BODY_PT, room, TITLE_RULE, ids, bold=True, keep=True)
    if "kinh_gui" in fields:
        _kinh_gui(document, text("kinh_gui"), room)

    # ----- legal basis and body -----
    basis = _lines(text("can_cu", ""))
    for index, line in enumerate(basis):
        ending = "." if index == len(basis) - 1 else ";"
        _line(document, line.rstrip(" .;,") + ending, BODY_PT, WD_ALIGN_PARAGRAPH.JUSTIFY, italic=bool(template.decision_word), before=6,
              first_cm=1.0)
    if template.decision_word:
        _line(document, template.decision_word, BODY_PT, bold=True, before=6)
    body, appendices = split(content.blocks)
    articles = any(_LEAD.match(b["text"]) for b in content.blocks if b["kind"] in ("heading", "para"))
    _blocks(document, body, room, articles=articles, close=True)

    # ----- recipients | signature -----
    footer = document.add_table(rows=1, cols=2)
    footer.alignment = WD_TABLE_ALIGNMENT.CENTER
    _borderless(footer)
    _columns(footer, FOOT_CM)
    recipients, signature = footer.rows[0].cells
    _line(recipients, "Nơi nhận:", SMALL_PT, WD_ALIGN_PARAGRAPH.LEFT, bold=True, italic=True, before=12, reuse=recipients.paragraphs[0])
    for line in _lines(text("noi_nhan", "- Lưu: VT, …. ….")):
        _line(recipients, line, RECIPIENT_PT, WD_ALIGN_PARAGRAPH.LEFT)
    authority = _lines(text("quyen_han")) or ["…"]
    for index, line in enumerate(authority):
        _line(signature, line.upper(), BODY_PT, bold=True, before=12 if index == 0 else 0, reuse=signature.paragraphs[0] if index == 0 else None)
    for _ in range(4):
        _line(signature, "", BODY_PT)
    _line(signature, text("nguoi_ky"), BODY_PT, bold=True)

    # ----- appendices (Mẫu 15a) -----
    cited = _citation(content, text, issuer_name)
    for appendix in appendices:
        section = document.add_section(WD_SECTION.NEW_PAGE)
        room = _page(section, landscape=appendix.landscape)
        _restart_numbering(section)
        _line(document, appendix.label, BODY_PT, bold=True, keep=True)
        _line(document, appendix.block["text"].upper(), BODY_PT, bold=True, keep=True)
        _ruled(document, cited, BODY_PT, room, KEM_THEO_RULE, ids, italic=True, keep=True)
        _line(document, "", BODY_PT)
        _blocks(document, appendix.blocks, room, articles=articles, close=False)

    target.parent.mkdir(parents=True, exist_ok=True)
    document.save(target)
    return target


def _citation(content: DraftContent, text: Callable[..., str], issuer_name: str) -> str:
    """The line under an appendix title that ties it to the document.

    Args:
        content: The draft.
        text: The field reader of `render` (placeholders and check samples applied).
        issuer_name: The issuing body in sentence case; the header's name, capitalised, when empty.

    Returns:
        "(Kèm theo Công văn số …/TM-CNTT ngày … tháng … năm … của Bộ Tham mưu)".
    """
    number = re.sub(r"^\s*số\s*:?\s*", "", text("so_ky_hieu", "…/…"), flags=re.IGNORECASE)
    dated = re.search(r"ngày.*$", text("dia_danh_ngay", "")) if "dia_danh_ngay" in content.fields else None
    issuer = issuer_name or text("co_quan_ban_hanh", "").capitalize()
    parts = [f"Kèm theo {content.template.name} số {number}", dated.group(0) if dated else "ngày … tháng … năm …"]
    return f"({' '.join(parts)}{f' của {issuer}' if issuer and issuer != '…' else ''})"


def slug(title: str) -> str:
    """A file-name-safe form of a title.

    Args:
        title: The draft's title.

    Returns:
        The title without path-hostile characters (Vietnamese letters are kept), at most 80 characters, or "du-thao".
    """
    cleaned = re.sub(r'[\\/:*?"<>|\r\n]+', " ", title).strip()
    return cleaned[:80] or "du-thao"
