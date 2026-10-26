"""Drafting: templates from the TT68 matrix, ownership rules, the renderer's output checked by doccheck, and the draft service."""
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest import mock

from app.core.errors import Conflict, ValidationFailed
from app.drafting import style, templates
from app.drafting.check import check_draft
from app.drafting.content import DraftContent, Issuer
from app.drafting.render import render
from sqlalchemy import select

from app.models import ChatSession, DraftRevision, User
from app.services import drafts

from support import memory_session

ISSUER = Issuer("Phòng Tác chiến", "Bộ Tham mưu")
BODY = [{"kind": "para", "text": "Thực hiện kế hoạch huấn luyện năm 2026, Phòng Tác chiến triển khai tập huấn như sau:"},
        {"kind": "heading", "text": "1. Thời gian, địa điểm"},
        {"kind": "para", "text": "Từ ngày 10/6/2026 đến ngày 20/6/2026 tại Trường Quân sự Quân khu."}]


def filled(code="M08") -> DraftContent:
    draft = DraftContent.new(templates.get(code), ISSUER, date(2026, 10, 26))
    draft.set_fields({"trich_yeu": "tổ chức tập huấn công tác tác chiến năm 2026", "kinh_gui": "Các đơn vị trực thuộc",
                      "can_cu": "Căn cứ Kế hoạch số 12/KH-BTM ngày 01/01/2026;", "noi_nhan": "- Như trên;\n- Lưu: VT, TC. A5."}, "assistant")
    draft.set_fields({"quyen_han": "TRƯỞNG PHÒNG", "nguoi_ky": "Đại tá Nguyễn Văn A", "dia_danh_ngay": "Hải Phòng, ngày 26 tháng 10 năm 2026"}, "user")
    draft.set_content(BODY)
    return draft


class TemplateTest(unittest.TestCase):
    def test_every_template_follows_its_model(self):
        for code, template in templates.all_templates().items():
            self.assertIn("quoc_hieu", template.fields, code)
            self.assertIn("co_quan_ban_hanh", template.fields, code)
            self.assertLessEqual(template.required, set(template.fields), code)
        cv, kh = templates.get("CV"), templates.get("KH")
        self.assertIn("kinh_gui", cv.fields)
        self.assertNotIn("ten_loai", cv.fields)  # a Công văn has no document-type heading
        self.assertTrue(kh.title_word == "KẾ HOẠCH" and "kinh_gui" not in kh.fields)
        self.assertTrue(templates.get("QĐ").has_can_cu and templates.get("QĐ").decision_word)

    def test_kinds_sharing_a_model_have_distinct_codes(self):
        self.assertEqual(templates.get("BC").code, "M06.BC")
        self.assertEqual(templates.get("M06.TB").name, "Thông báo")
        with self.assertRaises(KeyError):
            templates.get("nope")


class OwnershipTest(unittest.TestCase):
    def setUp(self):
        self.draft = DraftContent.new(templates.get("CV"), ISSUER, date(2026, 10, 26))

    def test_new_draft_is_filled_by_the_system(self):
        fields = self.draft.fields
        self.assertEqual(fields["co_quan_ban_hanh"], {"text": "PHÒNG TÁC CHIẾN", "owner": "system"})
        self.assertEqual(fields["co_quan_chu_quan"]["text"], "BỘ THAM MƯU")
        self.assertEqual(fields["quoc_hieu"]["text"], "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM")
        self.assertIn("ngày 26 tháng 10 năm 2026", fields["dia_danh_ngay"]["text"])
        self.assertEqual(self.draft.missing(), ["dia_danh_ngay", "trich_yeu", "kinh_gui", "quyen_han", "nguoi_ky", "noi_nhan"])
        self.assertEqual(fields["noi_nhan"]["text"], "- Như trên;\n- Lưu: VT, …. ….")

    def test_assistant_cannot_touch_system_or_user_text_but_can_fill_blanks(self):
        applied, refused = self.draft.set_fields({"co_quan_ban_hanh": "X", "trich_yeu": "về việc A", "nguoi_ky": "Đại tá B", "nope": "x"}, "assistant")
        self.assertEqual(applied, ["trich_yeu", "nguoi_ky"])
        self.assertEqual(set(refused), {"co_quan_ban_hanh", "nope"})
        self.draft.set_fields({"nguoi_ky": "Đại tá C"}, "user")
        _, refused = self.draft.set_fields({"nguoi_ky": "Đại tá D"}, "assistant")
        self.assertIn("nguoi_ky", refused)
        self.assertEqual(self.draft.fields["nguoi_ky"]["text"], "Đại tá C")

    def test_user_may_override_system_fields(self):
        applied, _ = self.draft.set_fields({"co_quan_ban_hanh": "PHÒNG KHÁC"}, "user")
        self.assertEqual((applied, self.draft.fields["co_quan_ban_hanh"]["owner"]), (["co_quan_ban_hanh"], "user"))

    def test_set_content_is_refused_over_user_text_and_proposals_take_over(self):
        ids = self.draft.set_content(BODY)
        self.draft.user_edit(blocks=[{"id": ids[0], "text": "Đoạn do người dùng viết."}])
        with self.assertRaises(ValidationFailed):
            self.draft.set_content(BODY)
        self.draft.propose(ids[0], "replace", "Đề xuất khác.", "gọn hơn")
        self.assertEqual(self.draft.block(ids[0])["text"], "Đoạn do người dùng viết.")  # untouched until accepted
        self.draft.accept(ids[0])
        self.assertEqual(self.draft.block(ids[0])["text"], "Đề xuất khác.")

    def test_reject_and_insert_and_delete_proposals(self):
        ids = self.draft.set_content(BODY)
        self.draft.propose(ids[1], "delete")
        self.draft.reject(ids[1])
        self.assertEqual(len(self.draft.blocks), 3)
        added = self.draft.propose(ids[0], "insert", "Đoạn thêm.")["id"]
        self.assertEqual([b["id"] for b in self.draft.blocks][1], added)
        self.assertEqual(self.draft.block(added)["text"], "")  # shown as a pending insertion
        self.draft.reject(added)
        self.assertEqual(len(self.draft.blocks), 3)
        self.draft.propose(ids[1], "delete")
        self.draft.accept(ids[1])
        self.assertEqual([b["id"] for b in self.draft.blocks], [ids[0], ids[2]])

    def test_block_ids_are_never_reused(self):
        first = self.draft.set_content(BODY)
        second = self.draft.set_content(BODY)
        self.assertFalse(set(first) & set(second))


class RenderTest(unittest.TestCase):
    def test_a_well_formed_cong_van_has_no_format_errors(self):
        result = check_draft(filled("M08"))
        self.assertEqual(result["counts"]["error"], 0, result["findings"])
        self.assertEqual(result["counts"]["warning"], 0, result["findings"])

    def test_kinds_with_a_title_word_pass_too(self):
        for code in ("M07", "M06.BC", "M06.TB"):
            result = check_draft(filled(code))
            self.assertEqual(result["counts"]["error"], 0, (code, result["findings"]))

    def test_findings_name_what_is_wrong(self):
        draft = filled("M08")
        draft.set_fields({"noi_nhan": "- Lưu: VT."}, "user")
        messages = " ".join(f["message"] for f in check_draft(draft)["findings"])
        self.assertIn("Lưu", messages)

    def test_the_closing_mark_is_added_and_placeholders_stay_visible(self):
        from docx import Document
        with tempfile.TemporaryDirectory() as folder:
            draft = filled("M08")
            draft.fields["so_ky_hieu"]["text"] = "Số: …/CV-…"
            path = render(draft, Path(folder) / "a.docx")
            text = "\n".join(p.text for p in Document(str(path)).paragraphs) + "\n".join(
                c.text for t in Document(str(path)).tables for r in t.rows for c in r.cells)
        self.assertIn("quân khu./.", text.lower().replace("quân sự quân khu", "quân khu"))
        self.assertIn("Số: …/CV-…", text)


class DraftServiceTest(unittest.TestCase):
    def setUp(self):
        folder = tempfile.TemporaryDirectory()
        self.addCleanup(folder.cleanup)
        patch = mock.patch.object(style, "profile_path", return_value=Path(folder.name) / "profiles.json")  # no style profiles
        patch.start()
        self.addCleanup(patch.stop)
        self.db = memory_session()
        self.user = User(username="u", password_hash="x", full_name="U", unit_id="btm.tac_chien", role="staff")
        self.db.add(self.user)
        self.db.flush()
        self.session = ChatSession(id="s1", user_id=self.user.id)
        self.db.add(self.session)
        self.db.flush()

    def test_create_uses_the_users_unit_and_refuses_a_second_draft(self):
        row = drafts.create(self.db, self.user, self.session, "CV")
        view = drafts.view(row)
        # a phòng of Bộ Tham mưu issues under the agency's name (Thông tư 68: the office only shows in the symbol and "Lưu")
        self.assertEqual(view["fields"]["co_quan_ban_hanh"]["text"], "BỘ THAM MƯU")
        self.assertEqual(view["fields"]["co_quan_chu_quan"]["text"], "QUÂN KHU 3")
        self.assertEqual((view["version"], view["template"]["name"]), (1, "Công văn"))
        with self.assertRaises(Conflict):
            drafts.create(self.db, self.user, self.session, "KH")
        self.assertEqual(drafts.create(self.db, self.user, self.session, "KH", replace=True).template, "M07")
        with self.assertRaises(ValidationFailed) as ctx:
            drafts.create(self.db, self.user, self.session, "XYZ", replace=True)
        self.assertIn("CV (Công văn)", ctx.exception.message)

    def test_every_change_bumps_the_version_and_is_logged(self):
        row = drafts.create(self.db, self.user, self.session, "CV")
        drafts.set_fields(self.db, row, {"trich_yeu": "về việc A"})
        ids = drafts.set_content(self.db, row, BODY)
        drafts.propose(self.db, row, ids[0], "replace", "Mới.")
        drafts.decide(self.db, row, ids[0], accept=True)
        drafts.user_edit(self.db, row, fields={"nguoi_ky": "Đại tá A"})
        view = drafts.view(row)
        self.assertEqual(view["version"], 6)
        self.assertEqual([h["by"] for h in view["history"]], ["assistant", "assistant", "assistant", "assistant", "user", "user"])
        self.assertEqual(view["title"], "về việc A")

    def test_export_writes_a_docx(self):
        row = drafts.create(self.db, self.user, self.session, "CV")
        drafts.set_content(self.db, row, BODY)
        with tempfile.TemporaryDirectory() as folder:
            path = drafts.export(self.db, row, Path(folder))
            self.assertTrue(path.name.endswith(".docx") and path.stat().st_size > 5000)

    def test_a_style_choice_is_checked_and_logged(self):
        row = drafts.create(self.db, self.user, self.session, "CV")
        view = drafts.view(row)
        self.assertEqual(view["style"], {"preset": "don_vi", "detail": "vua"})
        self.assertTrue(view["style_options"]["presets"])
        drafts.set_style(self.db, row, {"preset": "không có", "detail": "ngan"})
        view = drafts.view(row)
        self.assertEqual(view["style"], {"preset": "don_vi", "detail": "ngan"})
        self.assertEqual(view["version"], 2)
        self.assertIn("Ngắn", view["history"][-1]["summary"])
        drafts.set_style(self.db, row, {"preset": "don_vi", "detail": "ngan"})  # unchanged: no new version
        self.assertEqual(drafts.view(row)["version"], 2)

    def test_a_decision_and_a_later_edit_are_recorded_and_removed_with_the_draft(self):
        row = drafts.create(self.db, self.user, self.session, "CV")
        ids = drafts.set_content(self.db, row, [{"kind": "para", "text": "Đoạn gốc."}])
        drafts.propose(self.db, row, ids[0], "replace", "Đảm bảo xong 18/20 việc.")
        checks = drafts.content_of(row).blocks[0]["proposal"]["checks"]
        self.assertTrue(any(item["message"] == "Có thêm '18/20' không có trong đoạn gốc." for item in checks))
        self.assertTrue(any("bảo đảm" in item["message"] for item in checks))
        drafts.decide(self.db, row, ids[0], accept=True)
        drafts.user_edit(self.db, row, blocks=[{"id": ids[0], "text": "Bảo đảm xong việc."}])
        drafts.user_edit(self.db, row, blocks=[{"id": ids[0], "text": "Bảo đảm xong phần việc được giao."}])  # the person's own text now
        recorded = list(self.db.scalars(select(DraftRevision).where(DraftRevision.session_id == "s1").order_by(DraftRevision.id)))
        self.assertEqual([item.outcome for item in recorded], ["accepted", "edited_after"])
        self.assertEqual((recorded[0].unit_id, recorded[0].before, recorded[0].final), ("btm.tac_chien", "Đoạn gốc.", "Đảm bảo xong 18/20 việc."))
        self.assertEqual(recorded[1].final, "Bảo đảm xong việc.")
        drafts.propose(self.db, row, ids[0], "replace", "Bản khác.")
        drafts.decide(self.db, row, ids[0], accept=False)
        self.assertEqual(self.db.scalar(select(DraftRevision.outcome).order_by(DraftRevision.id.desc())), "rejected")
        drafts.delete(self.db, "s1")
        self.assertEqual(list(self.db.scalars(select(DraftRevision))), [])

    def test_user_edit_with_a_refused_field_fails_loudly(self):
        row = drafts.create(self.db, self.user, self.session, "CV")
        with self.assertRaises(ValidationFailed):
            drafts.user_edit(self.db, row, fields={"khong_co": "x"})


if __name__ == "__main__":
    unittest.main()
