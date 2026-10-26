import shutil
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.parsing import ParseError, parse_file
from app.parsing.legacy import legacy_font, tcvn3_to_unicode
from app.parsing.structure import classify

C, J = WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.JUSTIFY

def add(container, text, align=None, bold=False):
    p = container.add_paragraph()
    run = p.add_run(text)
    run.bold = bold
    if align is not None:
        p.alignment = align
    return p

def cell_lines(cell, lines, align=C):
    cell.paragraphs[0]._element.getparent().remove(cell.paragraphs[0]._element)
    for text, bold in lines:
        add(cell, text, align, bold)

def build_decision(path: Path, number="15", day="01", annex=False, stamp="") -> None:
    doc = Document()
    if stamp:
        add(doc, stamp, WD_ALIGN_PARAGRAPH.RIGHT, True)
    head = doc.add_table(rows=1, cols=2)
    cell_lines(head.cell(0, 0), [("BỘ QUỐC PHÒNG", False), ("QUÂN KHU 3", True), (f"Số: {number}/QĐ-QK3", False)])
    cell_lines(head.cell(0, 1), [("CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", True), ("Độc lập - Tự do - Hạnh phúc", True),
                                 (f"Hải Phòng, ngày {day} tháng 10 năm 2026", False)])
    add(doc, "QUYẾT ĐỊNH", C, True)
    add(doc, "Về việc kiện toàn Ban Chỉ đạo", C, True)
    add(doc, "TƯ LỆNH QUÂN KHU 3", C, True)
    add(doc, "Căn cứ Luật Quốc phòng ngày 08/6/2005;", J)
    add(doc, "Theo đề nghị của Chánh Văn phòng Quân khu.", J)
    add(doc, "QUYẾT ĐỊNH:", C, True)
    add(doc, "Điều 1. Kiện toàn Ban Chỉ đạo.", J, True)
    add(doc, "1. Trưởng ban: Phó Tư lệnh Quân khu.", J)
    add(doc, "2. Các thành viên:", J)
    add(doc, "a) Chỉ huy Văn phòng;", J)
    add(doc, "b) Chỉ huy phòng Tác chiến.", J)
    add(doc, "- Kinh phí 10.000 đồng mỗi người.", J)
    add(doc, "Điều 2. Hiệu lực thi hành.", J, True)
    add(doc, "Quyết định này có hiệu lực từ ngày ký.", J)
    foot = doc.add_table(rows=1, cols=2)
    cell_lines(foot.cell(0, 0), [("Nơi nhận:", True), ("- Như Điều 2;", False), ("- Lưu: VT, TM.", False)], WD_ALIGN_PARAGRAPH.LEFT)
    cell_lines(foot.cell(0, 1), [("TƯ LỆNH", True), ("Trung tướng Nguyễn Văn An", True)])
    if annex:
        add(doc, "PHỤ LỤC I", C, True)
        add(doc, "Danh sách thành viên kèm theo Quyết định số 15/QĐ-QK3", C)
        add(doc, "1. Đồng chí Nguyễn Văn B, Chỉ huy Văn phòng, thành viên.", J)
        add(doc, "2. Đồng chí Trần Văn C, Chỉ huy phòng Tác chiến, thành viên.", J)
    doc.save(path)

class ParseDecisionTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = Path(tempfile.mkdtemp())
        cls.path = cls.tmp / "qd.docx"
        build_decision(cls.path)
        cls.doc = parse_file(cls.path)
        cls.meta = cls.doc.metadata

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def test_header_metadata(self):
        m = self.meta
        self.assertEqual((m.doc_type, m.doc_type_code), ("Quyết định", "QĐ"))
        self.assertEqual((m.number, m.symbol), ("15", "QĐ-QK3"))
        self.assertEqual((m.parent_agency, m.issuing_agency), ("BỘ QUỐC PHÒNG", "QUÂN KHU 3"))
        self.assertEqual((m.place, m.date), ("Hải Phòng", "2026-10-01"))
        self.assertEqual(m.title, "Về việc kiện toàn Ban Chỉ đạo")
        self.assertFalse(m.is_draft)

    def test_legal_bases(self):
        self.assertEqual(len(self.meta.legal_bases), 2)
        self.assertTrue(self.meta.legal_bases[0].startswith("Căn cứ Luật Quốc phòng"))

    def test_footer(self):
        m = self.meta
        self.assertEqual(m.distribution, ["Như Điều 2", "Lưu: VT, TM"])
        self.assertEqual((m.signer_role, m.signer_name), ("TƯ LỆNH", "Trung tướng Nguyễn Văn An"))
        self.assertEqual(self.doc.warnings, [])

    def test_section_tree(self):
        dieu = [c for c in self.doc.root.children if c.kind == "dieu"]
        self.assertEqual([d.marker for d in dieu], ["Điều 1", "Điều 2"])
        khoan = dieu[0].children
        self.assertEqual([k.marker for k in khoan], ["1", "2"])
        self.assertEqual([p.marker for p in khoan[1].children], ["a", "b"])
        self.assertEqual(khoan[1].children[1].blocks[0].kind, "list_item")  # "- Kinh phí 10.000 đồng" is not a numbered unit
        self.assertEqual(dieu[1].blocks[0].text, "Quyết định này có hiệu lực từ ngày ký.")

    def test_footer_is_not_in_the_body(self):
        self.assertNotIn("Nơi nhận", self.doc.text())
        self.assertNotIn("Nguyễn Văn An", self.doc.text())

    def test_annex_after_signature_is_kept_out_of_the_footer(self):
        path = self.tmp / "annex.docx"
        build_decision(path, annex=True)
        doc = parse_file(path)
        self.assertEqual(doc.metadata.signer_name, "Trung tướng Nguyễn Văn An")
        self.assertEqual(doc.metadata.distribution, ["Như Điều 2", "Lưu: VT, TM"])
        annex = [c for c in doc.root.children if c.kind == "annex"]
        self.assertEqual([a.heading for a in annex], ["PHỤ LỤC I"])
        self.assertIn("Trần Văn C", annex[0].full_text())
        self.assertNotIn("Trần Văn C", " ".join(doc.footer))

    def test_secrecy_stamp(self):
        self.assertEqual(self.doc.metadata.classification, "")
        for stamp, expected in [("MẬT", "mat"), ("TỐI MẬT", "toi_mat"), ("(TUYỆT MẬT)", "tuyet_mat")]:
            path = self.tmp / "s.docx"
            build_decision(path, stamp=stamp)
            self.assertEqual(parse_file(path).metadata.classification, expected, stamp)

    def test_draft_has_blank_number_and_date(self):
        path = self.tmp / "draft.docx"
        build_decision(path, number="", day="")
        m = parse_file(path).metadata
        self.assertTrue(m.is_draft)
        self.assertEqual((m.number, m.symbol, m.date), ("", "QĐ-QK3", ""))

class ClassifyTests(unittest.TestCase):
    def test_numbered_units(self):
        cases = {
            "Điều 12. Trách nhiệm": ("dieu", "Điều 12"), "I. MỤC ĐÍCH": ("roman", "I"), "IV. TỔ CHỨC": ("roman", "IV"),
            "1. Nội dung": ("number", "1"), "2) Thời gian": ("number", "2"), "a) Lực lượng": ("letter", "a"),
            "đ) Khác": ("letter", "đ"), "1.2. Kinh phí": ("decimal", "1.2"), "A. PHẦN MỞ ĐẦU": ("upper", "A"),
            "Chương II": ("chuong", "Chương II"), "Câu 3. Theo tài liệu": ("cau", "Câu 3"),
        }
        for text, (kind, marker) in cases.items():
            got = classify(text)
            self.assertIsNotNone(got, text)
            self.assertEqual((got[0], got[1]), (kind, marker), text)

    def test_not_numbered_units(self):
        for text in ["10.000 đồng mỗi người", "1.5 triệu đồng", "Mục đích, yêu cầu", "M. Nguyễn Văn A", "2026 là năm", "Căn cứ Luật"]:
            self.assertIsNone(classify(text), text)

class LegacyFontTests(unittest.TestCase):
    def test_font_detection(self):
        self.assertEqual(legacy_font(".VnTime"), "tcvn3")
        self.assertEqual(legacy_font(".VnArialH"), "tcvn3_upper")
        self.assertIsNone(legacy_font("Times New Roman"))
        self.assertIsNone(legacy_font(None))

    def test_tcvn3_conversion(self):
        self.assertEqual(tcvn3_to_unicode("Th«ng b¸o ch\xad¬ng tr×nh"), "Thông báo chương trình")
        self.assertEqual(tcvn3_to_unicode("Đã là Unicode ư"), "Đã là Unicode ư")  # words with non-Latin-1 chars are left alone

class ReadErrorTests(unittest.TestCase):
    def test_missing_and_unsupported(self):
        with self.assertRaises(ParseError):
            parse_file("/nonexistent/file.docx")
        tmp = Path(tempfile.mkdtemp())
        try:
            (tmp / "a.xyz").write_text("x")
            with self.assertRaises(ParseError):
                parse_file(tmp / "a.xyz")
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

    def test_misnamed_extension_is_read_by_content(self):
        tmp = Path(tempfile.mkdtemp())
        try:
            build_decision(tmp / "real.docx")
            shutil.copy(tmp / "real.docx", tmp / "wrong.doc")
            doc = parse_file(tmp / "wrong.doc")
            self.assertEqual(doc.format, "docx")
            self.assertTrue(any("không khớp nội dung" in w for w in doc.warnings))
        finally:
            shutil.rmtree(tmp, ignore_errors=True)

if __name__ == "__main__":
    unittest.main()
